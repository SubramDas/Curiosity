"""Build the editable hackathon deck; run with python-pptx, Pillow and CairoSVG."""
from pathlib import Path
import json, math
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from PIL import Image
import cairosvg
ROOT=Path(__file__).resolve().parent
D=json.loads((ROOT/'deck-details.json').read_text())
P=Presentation(); P.slide_width=Inches(13.3333);P.slide_height=Inches(7.5)
C={'bg':'F5F7F2','ink':'17221D','green':'176B4C','mint':'55BF95','muted':'627169','line':'DCE4DC','white':'FFFFFF','dark':'102B22','pale':'E6F0E8','gold':'E5BA6A'}
SPECS=[]
cairosvg.svg2png(url=str(ROOT.parent/'public/brand/revelink-symbol.svg'),write_to=str(ROOT/'assets/logo.png'),output_width=400,output_height=400)
def rgb(c):return RGBColor.from_string(C.get(c,c))
def rect(s,x,y,w,h,c='white',radius=False,line=None):
 sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h));sh.fill.solid();sh.fill.fore_color.rgb=rgb(c);sh.line.fill.background()
 sh._element.spPr.append(OxmlElement('a:effectLst'))
 if radius: sh.adjustments[0]=.12
 if line:sh.line.color.rgb=rgb(line)
 SPECS[-1]['shapes'].append(dict(type='rect',x=x,y=y,w=w,h=h,color=C.get(c,c)))
 return sh
def text(s,x,y,w,h,t,size=20,color='ink',bold=False,font='DejaVu Sans',align=None):
 sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=sh.text_frame;tf.word_wrap=True
 tf.margin_left=tf.margin_right=0;tf.margin_top=tf.margin_bottom=0
 for i,line in enumerate(t.split('\n')):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.name=font;p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=rgb(color);p.space_after=Pt(6)
  if align is not None:p.alignment=align
 SPECS[-1]['shapes'].append(dict(type='text',x=x,y=y,w=w,h=h,text=t,size=size,color=C.get(color,color),bold=bold))
 return sh
def line(s,x,y,x2,y2,color='mint'):
 sh=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x),Inches(y),Inches(x2),Inches(y2));sh.line.color.rgb=rgb(color);sh.line.width=Pt(2)
def pic(s,path,x,y,w,h):
 path=Path(path);im=Image.open(path);iw,ih=im.size;scale=min(w/iw,h/ih);ww,hh=iw*scale,ih*scale
 s.shapes.add_picture(str(path),Inches(x+(w-ww)/2),Inches(y+(h-hh)/2),width=Inches(ww),height=Inches(hh))
def slide(section,title,subtitle='',dark=False,source=''):
 s=P.slides.add_slide(P.slide_layouts[6]);SPECS.append({'name':section,'dark':dark,'shapes':[]})
 s.background.fill.solid();s.background.fill.fore_color.rgb=rgb('dark' if dark else 'bg')
 text(s,.55,.32,11,.3,section.upper(),11,'mint' if dark else 'green',True)
 text(s,.55,.88,12.1,1.12,title,31,'white' if dark else 'ink',True)
 if subtitle:text(s,.57,2.02,12,.58,subtitle,15,'pale' if dark else 'muted')
 line(s,.55,7.03,12.78,7.03,'green' if dark else 'line')
 text(s,.55,7.12,10.8,.2,source or 'REVELINK  /  SSRD  /  AI BUILDER CUP 2026',8,'pale' if dark else 'muted')
 text(s,12.1,7.1,.65,.25,f'{len(P.slides):02}',10,'mint' if dark else 'green',True,align=PP_ALIGN.RIGHT)
 return s
def note(s,t):s.notes_slide.notes_text_frame.text=t

def card(s,x,y,w,h,kicker,title,body,dark=False):
 rect(s,x,y,w,h,'green' if dark else 'white',True)
 text(s,x+.23,y+.2,w-.46,.3,kicker,11,'mint' if dark else 'green',True)
 text(s,x+.23,y+.68,w-.46,.8,title,22,'white' if dark else 'ink',True)
 text(s,x+.23,y+1.58,w-.46,h-1.7,body,16,'pale' if dark else 'muted')

s=slide('Team details · Media, Content & Digital Experiences','',dark=True)
pic(s,ROOT/'assets/logo.png',.55,.9,.73,.73)
text(s,1.45,.92,9,.7,'Revelink',40,'white',True)
text(s,.6,2.12,11.8,1.65,'Keep the curiosity.\nKeep the learning connected.',38,'white',True)
text(s,.63,4.05,10.9,.65,'Watch. Discover. Connect the dots.',23,'mint')
text(s,.63,5.13,11.9,.72,'Problem: learners lose context and momentum moving between\nvideo discovery, explanations, and saved learning.',19,'pale')
text(s,.63,6.24,11,.35,f"TEAM  {D['team_name']}      TEAM LEADER  {D['team_leader']}",14,'white',True)
note(s,'Opening: I wanted to understand GPUs. I found videos, searched for explanations, and asked follow-up questions in separate tools. Revelink brings that journey together. The registered theme is Media, Content & Digital Experiences. The public prototype currently retains the earlier name Curiosity; the repository contains the approved Revelink rebrand. Team details supplied by founder. Original template administrative download note and blank closing slides omitted; all substantive prompts are covered.')

s=slide('Brief about the idea','A question should start a learning journey.', 'For adult students and early-career professionals in India, starting with English-language discovery.')
card(s,.55,2.9,3.85,3.46,'THE STARTING MOMENT','“How do GPUs work?”','A learner finds a video, but the next question sends them to another tool.')
card(s,4.73,2.9,3.85,3.46,'THE FRICTION','Context gets lost.','Search, watch, ask, and remember are separate tasks. Unrelated feeds interrupt the journey.')
card(s,8.91,2.9,3.85,3.46,'THE IDEA','One connected workspace.','YouTube discovery + a Gemini teacher + organised topics + saved playback.')
note(s,'Source: STARTUP_IDEA.md latest agreed direction and founder GPU example; src/App.tsx and shared feature modules. This is a founder-observed problem and product hypothesis, not a surveyed prevalence claim. Computing has prepared checked source notes; broader search does not imply reviewed teaching content across every subject.')

s=slide('Opportunities · differentiation · USP','The learning context travels with the learner.', 'USP: video discovery, in-context explanation, and the next useful action in one continuous journey.')
rows=[('Find a useful starting point','Search for videos separately','Educational intent checks and learning-angle choices'),('Resolve the doubt','Explain the context again in chat','Teacher uses question, conversation and eligible video'),('Act on the explanation','Find the section or search again','Jump to an AI-matched section or open a real result'),('Return to the interest','Reconstruct the previous session','Organised subjects and saved playback position')]
for x,w,t in [(.55,2.75,'LEARNER NEED'),(3.45,3.65,'FRAGMENTED WORKFLOW'),(7.25,5.53,'REVELINK WORKFLOW')]:
 rect(s,x,2.83,w,.48,'green');text(s,x+.14,2.93,w-.2,.3,t,11,'white',True)
for i,row in enumerate(rows):
 y=3.42+i*.72
 for x,w,t in zip([.55,3.45,7.25],[2.75,3.65,5.53],row):
  rect(s,x,y,w,.64,'white');text(s,x+.14,y+.1,w-.28,.5,t,13,'ink',x==.55)
text(s,.6,6.52,12,.25,'Starting opportunity: adult self-directed learners exploring a concept, not a full course.',13,'green',True)
note(s,'Comparison is a workflow framing, not an exhaustive competitor feature audit or a claim that other products cannot offer these capabilities. Desktop educational guardrails and Universe navigation are local implementations, not verified public-deployment features. Differentiation derives from integration and intentional transitions, not proprietary foundation-model ownership. Demand and willingness to pay remain unvalidated.')

s=slide('List of features offered','Built around discovery, clarity, and continuity.')
features=[('01 / DISCOVER','Find your starting point','Search, video results, long videos and a short-video feed.'),('02 / UNDERSTAND','Ask beside the video','Follow-ups, simpler explanations and relevant video sections.'),('03 / CONTINUE','Resume your learning','Saved topics, watch history and last playback position.'),('04 / ORGANISE','Explore three Universes','Science & Technology; Commerce & Business; Arts & Humanities.'),('05 / STAY ON TASK','Guide learning intent','Clarify ambiguous intent; filter results using metadata.'),('06 / RESET','Take a bounded break','Optional topic-related clips after 12 active minutes; at most three.')]
for i,(k,t,b) in enumerate(features):
 x=.55+(i%3)*4.18;y=2.32+(i//3)*2.17
 rect(s,x,y,3.85,1.96,'white',True);text(s,x+.2,y+.16,3.45,.24,k,10,'green',True);text(s,x+.2,y+.58,3.45,.4,t,18,'ink',True);text(s,x+.2,y+1.08,3.45,.76,b,13,'muted')
text(s,.6,6.78,12,.17,'Latest local build shown. Public demo is an earlier revision; learning outcomes have not yet been measured.',9,'muted')
note(s,'Sources: shared/features/videos, chat, topics, progress, learning; desktop-ui/features/universes and history; backend/search. Break active time is measured by the browser; the server gate checks elapsed session time and is not proof of attention. Account progress code exists but wider live two-account sync/isolation review remains. Graph is intentionally omitted from active feature list because it defaults off.')

s=slide('Process flow diagram','From “I wonder…” to “I can keep going.”','Illustrative GPU journey • implemented paths; AI decisions can be uncertain.')
steps=[('01','SEARCH','“How do GPUs work?”'),('02','CHOOSE','Clarify intent if needed'),('03','WATCH','Open a learning video'),('04','ASK','“Why useful for AI?”'),('05','CONTINUE','Resume or explore next')]
for i,(n,t,b) in enumerate(steps):
 x=.55+i*2.49;rect(s,x,3.0,2.26,1.7,'white',True);text(s,x+.16,3.17,1.9,.32,n,16,'green',True);text(s,x+.16,3.66,1.9,.35,t,17,'ink',True);text(s,x+.16,4.12,1.96,.5,b,12,'muted')
 if i<4:text(s,x+2.27,3.62,.22,.4,'›',23,'green',True)
rect(s,3.05,5.12,7.23,1.1,'pale',True);text(s,3.27,5.28,6.8,.72,'Teacher decision\nMatching section → seek  •  No match → explanation + actual video suggestions',15,'green',True)
line(s,9.1,4.72,9.1,5.1);line(s,10.28,5.67,11.65,5.67);line(s,11.65,5.67,11.65,4.72)
note(s,'Source: backend/search/routes.mjs, backend/chat/teacher.mjs and gemini-teacher.mjs, shared/features/chat/useTutorChat.ts and src/App.tsx. The teacher can answer another subject without forcing a player change. Timestamp requires eligible analyzed video, related question, valid bounds and nonempty evidence; it remains AI-matched and may be approximate. Ineligible video uses a clearly labeled general explanation without a timestamp.')

s=slide('Wireframes / mock diagrams','An interface designed around the learning task.','Actual local React interface • simulated API responses and player in these captures.')
pic(s,ROOT/'assets/local-search-choices.png',.55,2.7,6.0,3.65)
pic(s,ROOT/'assets/local-teacher.png',6.8,2.7,5.98,3.65)
text(s,.65,6.5,5.8,.35,'01  Clarify the learning angle before searching.',13,'green',True)
text(s,6.9,6.5,5.8,.35,'02  Keep the teacher beside the video.',13,'green',True)
note(s,'Captured by the repository browser regression tests on 2026-10-18, not a live AI-answer demonstration. Penpot is the project design source; inspected New File 1 / Revelink — Brand explorations and curiosity-desktop-universe through the connected MCP. Native text/shapes used throughout this PowerPoint; application screenshots are the only flattened panels. Source captures: tests/browser/search.spec.ts and chat-readability.spec.ts.')

s=slide('Architecture diagram','Google Cloud runs the application. Gemini powers the learning decisions.','Latest repository architecture • browser calls the API; provider credentials stay on the server.')
rect(s,.55,2.83,2.4,2.0,'white',True);text(s,.76,3.05,1.98,.35,'REACT CLIENT',16,'green',True);text(s,.76,3.6,1.98,1.1,'Desktop / mobile\nShared features\nGuest storage',14)
rect(s,3.5,2.72,5.0,3.63,'pale',True);text(s,3.72,2.9,4.55,.35,'GOOGLE CLOUD RUN / NODE API',16,'green',True)
text(s,3.75,3.48,4.44,1.1,'Validation + access checks\nTeacher • search • classification\nShared Gemini priority queue',15)
rect(s,3.73,4.85,4.54,1.13,'white',True);text(s,3.92,5.02,4.15,.8,'Firestore: metadata cache + atomic budgets\nSecret Manager: provider credentials',14)
line(s,2.98,3.8,3.46,3.8);text(s,2.98,3.35,.5,.3,'API',9,'green')
for i,(title,body) in enumerate([('GEMINI DEVELOPER API','Teaching, intent and topic decisions'),('YOUTUBE + WIKIDATA','Video metadata + concept candidates'),('SUPABASE','Auth + private topics / playback')]):
 y=2.72+i*1.24;rect(s,9.05,y,3.73,1.03,'white',True);text(s,9.23,y+.14,3.36,.28,title,12,'green',True);text(s,9.23,y+.55,3.35,.35,body,11);line(s,8.52,y+.5,9.02,y+.5)
text(s,.6,6.57,12,.27,'Video streams directly through the official YouTube player. No media is hosted by Revelink.',12,'muted')
note(s,'Sources: backend/production.mjs, backend/server.mjs, backend/storage, backend/http, backend/providers, backend/supabase.mjs, Dockerfile, deployment docs. Gemini uses Developer API REST, not Vertex AI. Google Cloud is primary app hosting; Supabase is supporting auth/private storage. The browser also uses the Supabase public auth client. Latest classification/search queue is local; the hosted Cloud Run revision predates those modules. One process-local Gemini queue requires distributed coordination before multiple replicas.')

s=slide('Technical depth · meaningful generative AI','Three decisions that connect the experience.',dark=True)
for i,(k,t,b) in enumerate([('TEACHER','Explain → act','Question + history + eligible video → structured answer, source basis and validated section or real video suggestions.'),('SEARCH','Intent → useful results','Question → learning choices or approval → YouTube retrieval → batched educational-purpose decisions.'),('CLASSIFICATION','Concept → discovery','Video metadata → Gemini topic extraction → Wikidata candidates → validated primary and related Universes.')]):
 card(s,.55+i*4.18,2.65,3.85,3.6,k,t,b,True)
text(s,.65,6.52,12,.3,'Validated outputs  /  explicit uncertainty  /  cached decisions  /  teacher-first scheduling',14,'mint',True)
note(s,'Source implementation: backend/chat/gemini-teacher.mjs validates structured answers and known source IDs; backend/chat/teacher.mjs retrieves actual YouTube recommendations. backend/search validates model decisions, IDs and intent. backend/classification performs two-stage extraction/resolution with Wikidata and provisional fallback. Sources do not imply factual verification of all videos. Gemini model is configured via GEMINI_MODEL; historical project notes record gemini-3.5-flash-lite, but this slide deliberately does not promise a particular billing/model tier. Queue priority teacher, search, background classification; default 15-second start spacing. No universal latency claim.')

s=slide('Technologies used','A small, modular stack with clear responsibilities.')
rows=[('Experience','React · TypeScript · Vite','Shared learning features; dedicated desktop and mobile layouts.'),('Application','Node.js · REST API','Provider adapters, validation, background jobs and budgets.'),('AI & discovery','Gemini API · YouTube Data API · Wikidata','Teaching, educational intent, video retrieval and concept labels.'),('Cloud platform','Cloud Run · Firestore · Secret Manager','App hosting, persistent cache/counters and managed secrets.'),('Accounts','Supabase Auth · PostgreSQL / RLS','Sign-in, user-owned topics and playback records.'),('Verification','Playwright · Node tests · Penpot','Repeatable browser journeys, service rules and editable designs.')]
for i,(a,b,c) in enumerate(rows):
 y=2.45+i*.67;rect(s,.55,y,12.23,.58,'white',True);text(s,.75,y+.13,1.65,.3,a,13,'green',True);text(s,2.55,y+.13,4.35,.3,b,13,'ink',True);text(s,7.02,y+.13,5.5,.35,c,12,'muted')
note(s,'Source: package.json, ARCHITECTURE.md, backend/production.mjs, migrations and browser tests. Production backend does not depend on experiments; shared UI does not import layout folders or backend. src coordinates feature state. Penpot approved Orbit Book visual language is reused. Technologies listed are present in repository, not prospective additions.')

s=slide('Estimated implementation cost · optional','Designed for a small, quota-controlled pilot.','Illustrative compute estimate, not an actual bill or a quote for the Mumbai deployment.')
rect(s,.55,2.78,4.1,3.55,'green',True);text(s,.83,3.05,3.55,.4,'REFERENCE COMPUTE',12,'mint',True);text(s,.83,3.72,3.55,.8,'≈ $2.74',40,'white',True);text(s,.83,4.72,3.55,1.15,'per month before free-tier credits\n30 active hours · 1 vCPU · 0.5 GiB\n30,000 requests',15,'white')
text(s,5.05,2.95,7.5,.5,'What the estimate includes',22,'ink',True)
text(s,5.05,3.62,7.45,1.05,'Reference request-based Cloud Run rates:\n108,000 sec × ($0.000024 + 0.5 × $0.0000025)\n+ 0.03 million requests × $0.40 = $2.739',16,'muted')
text(s,5.05,4.92,7.45,.4,'What must be priced separately',20,'ink',True)
text(s,5.05,5.48,7.45,1.03,'Gemini tokens / video input, Firestore, secrets, builds, image storage, network transfer, Supabase and labour. Region and billing tier change totals.',15,'muted')
note(s,'Estimate assumptions only, not observed demand. Cloud Run default us-central1 active request-based list prices from https://cloud.google.com/run/pricing, checked 2026-10-18: CPU .000024 USD/vCPU-sec, RAM .0000025 USD/GiB-sec, .40 USD/million requests. Mumbai may have different rates; compute estimate is a reference calculation, NOT total project cost. No free-tier entitlement assumed. Gemini model/tier must be priced at https://ai.google.dev/gemini-api/docs/pricing against actual billed input/output/video tokens. Prototype uses bounded daily allowances and caching; minimum instances zero, max one. No fabricated total spend. Founder labour excluded.')
text(s,.6,6.7,12,.19,'Rate reference: cloud.google.com/run/pricing  •  Full assumptions and exclusions are in the speaker notes.',9,'muted')

s=slide('Snapshots of the prototype','A deployed product judges can open today.','Live Cloud Run screenshots • 18 October 2026 • deployed under the earlier Curiosity name.')
pic(s,ROOT/'assets/hosted-search.png',.55,2.65,8.6,3.9)
pic(s,ROOT/'assets/hosted-mobile.png',9.5,2.55,2.9,4.02)
text(s,.67,6.62,8.3,.25,'Desktop: GPU results and the teacher workspace.',13,'green',True)
text(s,9.4,6.62,3.25,.25,'Same URL on a phone.',13,'green',True)
note(s,'Actual public site captures from https://curiosity-u2ubbw6fra-el.a.run.app taken 2026-10-18, not invented mockups. Homepage and GPU results were visible; this screenshot does not by itself prove chat answer quality or live playback. Newer Revelink branding, desktop guardrails, Universe classification and history refinements are local, not claimed deployed. YouTube thumbnail/content rights remain with creators. The phone capture uses a desktop browser mobile viewport, not a physical phone test.')

s=slide('Prototype performance report / benchmarking','Evidence you can inspect. Limits we can name.','Fresh local verification + one live hosted teacher request + historical retrieval evidence.')
for i,(n,t,desc) in enumerate([('95 / 95','Service / rule checks','10 experiment + 85 current tests'),('63 / 64','Browser journeys','Simulated APIs, accounts and player'),('PASS','Production build','Type checks + browser bundle')]):
 x=.55+i*4.18;rect(s,x,2.75,3.85,1.75,'white',True);text(s,x+.22,2.92,3.4,.64,n,30,'green',True);text(s,x+.22,3.64,3.4,.35,t,16,'ink',True);text(s,x+.22,4.09,3.4,.25,desc,10,'muted')
text(s,.65,4.9,12,.4,'Live hosted check · 18 October',20,'ink',True)
text(s,.65,5.48,12,.7,'2.53 seconds: one hosted teacher response + two video suggestions.\nEarlier YouTube check: 18 unique videos in 867 ms (13 October).',17,'muted')
text(s,.65,6.43,12,.4,'Single observations, not load tests. One legacy saved-search browser path failed; learning impact is unmeasured.',12,'green',True)
note(s,'Fresh run 2026-10-18: npm test 10+85=95 passed, npm run build passed (main JS bundle size advisory remains), npm run test:browser 63 passed / 1 failed. Failure: universes.spec.ts:228, Other saved searches is absent; do not interpret as a fully green browser suite. No production code changed as part of deck preparation. Logs in submission/evidence. Historical timings are from PHASE_2_RESULTS.md Successful live retry, 2026-10-13; they do not measure current queue, video analysis, load or p95 latency. One live hosted general question on 2026-10-18 returned HTTP 200 in 2529 ms including two actual YouTube suggestions. No video context or timestamp was requested in that run. No measured learning improvement, retention or user adoption is claimed.')

s=slide('Additional details / future development','Prove the learning value. Then expand the map.','A focused next step, with outcomes measured before broader scale.')
card(s,.55,2.75,3.85,3.7,'NEXT / RELEASE READINESS','Align the demo','Deploy the approved build, resolve the legacy-search regression, review real account sync and run a hosted end-to-end check.')
card(s,4.73,2.75,3.85,3.7,'NEXT / LEARNER PILOT','Test the core hypothesis','Recruit 10–15 adult learners. Compare time to resolve a doubt, task completion, context switches and voluntary return.')
card(s,8.91,2.75,3.85,3.7,'LATER / DISCOVERY','Connect adjacent ideas','Activate the reviewed learning graph, test progress assessments, and add distributed scheduling before multiple replicas.')
note(s,'Roadmap and pilot are proposals, not completed milestones. Graph code exists locally but defaults off: 30 draft concepts and 26 connections; migration/publication and real access checks outstanding; no published videos/events. Graph exploration is activity, not mastery. Monetisation is unvalidated: begin free, test willingness to pay only after learner value is demonstrated. Paid-course integrations require separate rights/partnerships and are deferred. Scaling requires more than Cloud Run replicas because Gemini scheduler is process-local. Main hypothesis: reduced friction while resolving a self-directed doubt; test against current learner workflow.')

s=slide('Submission links','Try one curiosity. Follow it through.',dark=True)
text(s,.62,2.3,11.9,.6,'Search “GPUs” → ask a question → explore → return later.',22,'mint',True)
for y,k,v,url in [(3.3,'FINAL PRODUCT',D['product_url'],D['product_url']),(4.35,'PUBLIC GITHUB REPOSITORY',D['repository_url'],D['repository_url']),(5.4,'DEMO VIDEO · UNDER 3 MINUTES',D['demo_video_url'],D['demo_video_url'] if D['demo_video_url'].startswith('https://') else None)]:
 text(s,.65,y,11.8,.25,k,10,'mint',True);sh=text(s,.65,y+.36,11.8,.45,v,19,'white')
 if url:sh.text_frame.paragraphs[0].runs[0].hyperlink.address=url
text(s,.65,6.55,11.8,.3,'SSRD  /  Subram Das  /  Watch. Discover. Connect the dots.',13,'pale')
note(s,'Public repository verified via unauthenticated GitHub API 2026-10-18. Hosted URL returned HTTP 200 and rendered video results. Founder will add demo video later; replace the marked field before submitting and verify public unauthenticated access. Official submission requirements: https://aibuildercup.com/themes.html and https://aibuildercup.com/Faqs.html. PDF export included because official requirements ask for deck/PPT converted to PDF. No deployment or repository publication performed here.')

# Keep hyperlinks readable on the dark closing slide in PDF and PowerPoint.
for rel in P.part.rels.values():
 if rel.reltype.endswith('/theme'):
  from lxml import etree
  root=etree.fromstring(rel.target_part.blob)
  ns={'a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
  for tag in ('hlink','folHlink'):
   node=root.find('.//a:clrScheme/a:'+tag,ns)
   if node is not None:
    for child in list(node):node.remove(child)
    col=OxmlElement('a:srgbClr');col.set('val','55BF95');node.append(col)
  rel.target_part._blob=etree.tostring(root)
P.core_properties.title='Revelink — AI Builder Cup 2026'
P.core_properties.subject='SSRD | Media, Content & Digital Experiences'
P.core_properties.author='SSRD — Subram Das'
P.core_properties.keywords='Revelink, Gemini, Google Cloud, learning, AI Builder Cup'
out=ROOT/'Revelink_AI_Builder_Cup_SSRD.pptx';P.save(out)
print(out)
print('Slides:',len(P.slides))
